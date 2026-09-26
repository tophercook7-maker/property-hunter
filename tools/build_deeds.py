"""Publish the recorded deed reference for every Garland County parcel.

Why this matters: the assessor's public portal (actDataScout) carries deed
history, and it returns 403 to automation, so Property Hunter has always had
to send people there by hand. The county namelist carries Book and Page for
94.6% of parcels. That is the pointer to the recorded instrument -- not the
instrument, and not a title opinion -- and publishing it closes the one gap a
working agent named first when asked what was missing.

Keyed by normalized parcel id so lookup.html can resolve it in the browser.
"""
from __future__ import annotations

import argparse, json, os, re, sys
from datetime import datetime, timezone

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "data", "garland_deeds.json")


def norm(v):
    return re.sub(r"[^0-9A-Za-z]", "", str(v or "")).upper()


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    a = ap.parse_args()
    import openpyxl
    wb = openpyxl.load_workbook(a.workbook, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(it)]

    deeds, subs, n = {}, {}, 0
    for r in it:
        row = dict(zip(hdr, r))
        n += 1
        pid = norm(row.get("PARCEL"))
        if not pid:
            continue
        book, page = row.get("Book"), row.get("Page")
        sd = (str(row.get("SubdDesc") or "").strip() or None)
        if book and page:
            deeds[pid] = f"{str(book).strip()}/{str(page).strip()}"
        if sd:
            subs[pid] = sd

    doc = {
        "built_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "county": "Garland", "county_fips": "05051",
        "source": "Garland County Assessor revised namelist 26-1543",
        "note": ("Book/Page is the county's pointer to the recorded deed at the Circuit Clerk. "
                 "It is not the instrument, not a chain of title, and not a title opinion. "
                 "The most recent deed on the assessor's roll can lag an unrecorded or newly "
                 "recorded transfer."),
        "format": "parcel id (letters and digits only) -> \"book/page\"",
        "deeds": deeds, "subdivisions": subs,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump(doc, open(OUT, "w"), separators=(",", ":"))
    print(f"rows {n:,} | deed refs {len(deeds):,} ({len(deeds)/n*100:.1f}%) | "
          f"subdivision names {len(subs):,} | {os.path.getsize(OUT)/1e6:.1f} MB -> {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
